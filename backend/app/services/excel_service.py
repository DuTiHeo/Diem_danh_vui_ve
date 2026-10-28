import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from io import BytesIO
import os
try:
    import xlrd
except ImportError:
    xlrd = None

class ExcelService:
    @staticmethod
    def parse_lms_excel(file_bytes: bytes, filename: str = "") -> dict:
        """
        Reads an Excel file from NEU LMS (.xlsx or .xls) across ALL sheets (can be 1 to 4+ sheets)
        and extracts metadata and the complete list of students.
        """
        is_xls = file_bytes.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1") or filename.lower().endswith(".xls")
        
        # List of sheets, where each sheet is a list of rows (each row is a list of strings)
        sheets_data: list[list[list[str]]] = []

        if is_xls:
            if not xlrd:
                raise ValueError("xlrd library is required to parse .xls files")
            wb = xlrd.open_workbook(file_contents=file_bytes)
            for s_idx in range(wb.nsheets):
                sheet = wb.sheet_by_index(s_idx)
                sheet_rows = []
                for r in range(sheet.nrows):
                    sheet_rows.append([str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)])
                if sheet_rows:
                    sheets_data.append(sheet_rows)
        else:
            wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), data_only=True)
            for sheet in wb.worksheets:
                sheet_rows = []
                for row in sheet.iter_rows(values_only=True):
                    sheet_rows.append([str(c).strip() if c is not None else "" for c in row])
                if sheet_rows:
                    sheets_data.append(sheet_rows)

        course_name = ""
        room_name = ""
        class_section = ""
        students = []
        seen_student_codes = set()

        for sheet_rows in sheets_data:
            header_idx = None
            ma_sv_col = -1
            ho_ten_col = -1
            ten_col = -1

            for idx, row in enumerate(sheet_rows):
                # Scan for metadata
                for cell in row:
                    if not course_name and ("Học phần:" in cell or "Môn học:" in cell):
                        course_name = cell.split(":", 1)[-1].strip()
                    if not room_name:
                        if "Phòng:" in cell or "Địa điểm:" in cell:
                            room_name = cell.split(":", 1)[-1].strip()
                        elif any(g in cell for g in ["Giảng đường", "B-2", "B-3", "A2-", "D2-"]):
                            room_name = cell.strip()
                    if not class_section:
                        if "Lớp học phần:" in cell:
                            class_section = cell.split(":", 1)[-1].strip()
                        elif "(" in cell and ")" in cell and len(cell) < 25:
                            class_section = cell.strip()

                # Find student header row for THIS sheet
                for c_idx, cell in enumerate(row):
                    if cell in ["Mã số SV", "Mã SV", "Mã sinh viên", "MSV"]:
                        header_idx = idx
                        ma_sv_col = c_idx
                        for n_idx, n_cell in enumerate(row):
                            if n_cell in ["Họ và tên", "Họ tên"]:
                                ho_ten_col = n_idx
                            elif n_cell in ["Họ và tên đệm", "Họ đệm"]:
                                ho_ten_col = n_idx
                            elif n_cell == "Tên":
                                ten_col = n_idx
                        break

                # Read student records
                if header_idx is not None and idx > header_idx:
                    if ma_sv_col != -1 and ma_sv_col < len(row):
                        st_code = row[ma_sv_col].replace(".0", "").strip()
                        if st_code.isdigit() and len(st_code) >= 6:
                            if st_code not in seen_student_codes:
                                seen_student_codes.add(st_code)
                                name = ""
                                if ho_ten_col != -1 and ho_ten_col < len(row):
                                    name = row[ho_ten_col]
                                if ten_col != -1 and ten_col < len(row):
                                    name = f"{name} {row[ten_col]}".strip()
                                students.append({
                                    "ordinal": len(students) + 1,
                                    "student_code": st_code,
                                    "full_name": name
                                })

        return {
            "course_name": course_name or "Lớp học phần",
            "room_name": room_name or "Phòng học",
            "class_section": class_section or "Lớp HP",
            "students": students,
            "total_sessions": 10
        }

    @staticmethod
    def save_template(file_bytes: bytes, course_id: str, upload_dir: str = "uploads") -> str:
        os.makedirs(upload_dir, exist_ok=True)
        file_path = os.path.join(upload_dir, f"{course_id}_template.xlsx")
        with open(file_path, "wb") as f:
            f.write(file_bytes)
        return file_path

    @staticmethod
    def generate_attendance_report(
        course_name: str,
        class_section: str,
        room_name: str,
        students: list[dict],
        sessions: list[dict],
        attendance_map: dict
    ) -> BytesIO:
        """
        Tạo file Excel kết quả điểm danh chuẩn NEU.
        students: [{"id": str, "ordinal": int, "student_code": str, "full_name": str}]
        sessions: [{"id": str, "session_number": int, "status": str}]
        attendance_map: {(student_id, session_id): "ATTENDED" | "ABSENT"}
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Diem_Danh"

        # Fonts & Styles
        title_font = Font(name="Calibri", size=15, bold=True, color="1E3A8A")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        bold_font = Font(name="Calibri", size=10, bold=True)
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        attended_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        absent_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

        thin_border = Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC")
        )

        # Meta headers
        ws.merge_cells("A1:G1")
        ws["A1"] = f"BẢNG ĐIỂM DANH SINH VIÊN - {course_name.upper()}"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        ws["A2"] = f"Lớp học phần: {class_section} | Phòng: {room_name}"
        ws["A2"].font = bold_font

        # Table headers
        headers = ["STT", "Mã SV", "Họ và tên"]
        for s in sessions:
            headers.append(f"Buổi {s['session_number']}")
        headers.extend(["Tổng có mặt", "Tổng vắng", "Tỷ lệ (%)"])

        start_row = 4
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Student rows
        num_sessions = len(sessions)
        for row_idx, st in enumerate(students, start=start_row + 1):
            ws.cell(row=row_idx, column=1, value=st.get("ordinal", row_idx - start_row)).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=2, value=st.get("student_code", "")).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=3, value=st.get("full_name", "")).alignment = Alignment(horizontal="left")

            attended_count = 0
            absent_count = 0

            for s_idx, session in enumerate(sessions, start=4):
                status = attendance_map.get((st["id"], session["id"]))
                cell = ws.cell(row=row_idx, column=s_idx)
                cell.alignment = Alignment(horizontal="center")
                cell.border = thin_border

                if status == "ATTENDED":
                    cell.value = "✓"
                    cell.fill = attended_fill
                    cell.font = bold_font
                    attended_count += 1
                elif status == "ABSENT":
                    cell.value = "V"
                    cell.fill = absent_fill
                    cell.font = bold_font
                    absent_count += 1
                else:
                    cell.value = "-"

            # Stats columns
            col_att = start_row + num_sessions + 1
            col_abs = col_att + 1
            col_pct = col_abs + 1

            c_att = ws.cell(row=row_idx, column=4 + num_sessions, value=attended_count)
            c_att.alignment = Alignment(horizontal="center")
            c_att.font = bold_font

            c_abs = ws.cell(row=row_idx, column=5 + num_sessions, value=absent_count)
            c_abs.alignment = Alignment(horizontal="center")
            c_abs.font = bold_font

            pct = round((attended_count / num_sessions * 100), 1) if num_sessions > 0 else 0
            c_pct = ws.cell(row=row_idx, column=6 + num_sessions, value=f"{pct}%")
            c_pct.alignment = Alignment(horizontal="center")
            c_pct.font = bold_font

            for c in range(1, 7 + num_sessions):
                ws.cell(row=row_idx, column=c).border = thin_border

        # Adjust column widths
        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 26
        for c in range(4, 7 + num_sessions):
            ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 12

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output


